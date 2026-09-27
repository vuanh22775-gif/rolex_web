from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_builder.common import add_body_p, add_bullet_p, add_heading_1, add_heading_2, add_table, add_table_title

def render_front_matter(doc):
    # ── 1. BÌA CHÍNH (Phụ lục 1) ──────────────────────────────────
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("TRƯỜNG ĐẠI HỌC PHƯƠNG ĐÔNG\nKHOA CÔNG NGHỆ SỐ & TRUYỀN THÔNG\nNGÀNH CÔNG NGHỆ THÔNG TIN\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(60)
    p_line.paragraph_format.first_line_indent = Cm(0)
    r_line = p_line.add_run("____________________\n\n\n\n\n")
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(14)
    r_line.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.first_line_indent = Cm(0)
    r_main = p_title.add_run("BÁO CÁO ĐỒ ÁN KỲ\n\n")
    r_main.font.name = 'Times New Roman'
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    
    r_sub = p_title.add_run("ĐỀ TÀI:\nỨNG DỤNG CÔNG NGHỆ WEB VÀO KINH DOANH ĐỒNG HỒ ROLEX TRỰC TUYẾN\n")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(16)
    r_sub.font.bold = True

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(160)
    p_foot.paragraph_format.space_after = Pt(0)
    p_foot.paragraph_format.first_line_indent = Cm(0)
    r_foot = p_foot.add_run("Hà Nội, 2026")
    r_foot.font.name = 'Times New Roman'
    r_foot.font.size = Pt(14)
    r_foot.font.bold = True

    doc.add_page_break()

    # ── 2. BÌA PHỤ (Phụ lục 1) ───────────────────────────────────
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(20)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.first_line_indent = Cm(0)
    r2 = p2.add_run("TRƯỜNG ĐẠI HỌC PHƯƠNG ĐÔNG\nKHOA CÔNG NGHỆ SỐ & TRUYỀN THÔNG\nNGÀNH CÔNG NGHỆ THÔNG TIN\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    r2.font.bold = True
    
    p2_line = doc.add_paragraph()
    p2_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_line.paragraph_format.space_before = Pt(0)
    p2_line.paragraph_format.space_after = Pt(40)
    p2_line.paragraph_format.first_line_indent = Cm(0)
    r2_line = p2_line.add_run("____________________\n\n\n")
    r2_line.font.name = 'Times New Roman'
    r2_line.font.size = Pt(14)
    r2_line.font.bold = True

    p2_title = doc.add_paragraph()
    p2_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_title.paragraph_format.space_before = Pt(10)
    p2_title.paragraph_format.space_after = Pt(14)
    p2_title.paragraph_format.first_line_indent = Cm(0)
    r2_main = p2_title.add_run("BÁO CÁO ĐỒ ÁN KỲ\n\n")
    r2_main.font.name = 'Times New Roman'
    r2_main.font.size = Pt(18)
    r2_main.font.bold = True
    
    r2_sub = p2_title.add_run("ĐỀ TÀI:\nỨNG DỤNG CÔNG NGHỆ WEB VÀO KINH DOANH ĐỒNG HỒ ROLEX TRỰC TUYẾN\n\n\n")
    r2_sub.font.name = 'Times New Roman'
    r2_sub.font.size = Pt(15)
    r2_sub.font.bold = True

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.space_before = Pt(30)
    p_info.paragraph_format.space_after = Pt(6)
    p_info.paragraph_format.left_indent = Cm(2.5)
    p_info.paragraph_format.line_spacing = 1.4
    p_info.paragraph_format.first_line_indent = Cm(0)
    
    info_text = (
        "Họ và tên sinh viên:\tVũ Nhật Tuấn Anh\n"
        "Mã sinh viên:\t\t523100132\n"
        "Lớp:\t\t\t523100C\n"
        "Ngành:\t\t\tCông nghệ thông tin\n\n"
        "Giảng viên hướng dẫn:\t[CẦN BỔ SUNG: TS/ThS. Tên Giảng viên hướng dẫn]\n"
    )
    r_info = p_info.add_run(info_text)
    r_info.font.name = 'Times New Roman'
    r_info.font.size = Pt(14)

    p2_foot = doc.add_paragraph()
    p2_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_foot.paragraph_format.space_before = Pt(80)
    p2_foot.paragraph_format.space_after = Pt(0)
    p2_foot.paragraph_format.first_line_indent = Cm(0)
    r2_foot = p2_foot.add_run("Hà Nội, 2026")
    r2_foot.font.name = 'Times New Roman'
    r2_foot.font.size = Pt(14)
    r2_foot.font.bold = True

    doc.add_page_break()

    # ── 3. MỤC LỤC ───────────────────────────────────────────────
    p_toc_head = doc.add_paragraph()
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_before = Pt(10)
    p_toc_head.paragraph_format.space_after = Pt(14)
    p_toc_head.paragraph_format.first_line_indent = Cm(0)
    r_toc = p_toc_head.add_run("MỤC LỤC")
    r_toc.font.name = 'Times New Roman'
    r_toc.font.size = Pt(16)
    r_toc.font.bold = True

    toc_entries = [
        ("MỞ ĐẦU", "5"),
        ("CHƯƠNG 1. TỔNG QUAN ĐỀ TÀI VÀ CƠ SỞ THỰC HIỆN", "8"),
        ("  1.1. Bài toán và lý do lựa chọn đề tài", "8"),
        ("  1.2. Mục tiêu của đề tài", "12"),
        ("  1.3. Đối tượng sử dụng, phạm vi và giới hạn của đề tài", "14"),
        ("  1.4. Dữ liệu đầu vào, đầu ra và các yêu cầu/ràng buộc chính", "16"),
        ("  1.5. Cơ sở lý thuyết, công nghệ và công cụ sử dụng", "19"),
        ("CHƯƠNG 2. PHÂN TÍCH VÀ THIẾT KẾ GIẢI PHÁP", "26"),
        ("  2.1. Phân tích yêu cầu và chức năng cốt lõi", "26"),
        ("  2.2. Use case hoặc các luồng xử lý chính", "30"),
        ("  2.3. Kiến trúc hệ thống, các thành phần và luồng trao đổi dữ liệu", "38"),
        ("  2.4. Thiết kế cơ sở dữ liệu", "42"),
        ("  2.5. Thiết kế giao diện, API và các thành phần kỹ thuật cần thiết", "47"),
        ("CHƯƠNG 3. XÂY DỰNG SẢN PHẨM, KIỂM THỬ VÀ ĐÁNH GIÁ", "53"),
        ("  3.1. Xây dựng và triển khai sản phẩm", "53"),
        ("  3.2. Kiểm thử và đánh giá sản phẩm", "61"),
        ("  3.3. Giải pháp cải tiến kỹ thuật", "70"),
        ("  3.4. Quá trình tự học và phát triển", "75"),
        ("  3.5. Kết quả đạt được và hạn chế", "79"),
        ("KẾT LUẬN", "82"),
        ("TÀI LIỆU THAM KHẢO", "84"),
        ("PHỤ LỤC", "86"),
        ("  Phụ lục 1: Hướng dẫn cài đặt và cấu hình hệ thống", "86"),
        ("  Phụ lục 2: Bảng chi tiết toàn bộ các Test Case kiểm thử hệ thống", "88"),
        ("  Phụ lục 3: Danh mục đầy đủ các RESTful API và đặc tả tham số", "92"),
    ]
    
    for title, page_num in toc_entries:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.line_spacing = 1.2
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(2)
        p_t.paragraph_format.first_line_indent = Cm(0)
        
        is_bold = not title.startswith("  ")
        r_title = p_t.add_run(title)
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(13)
        r_title.font.bold = is_bold
        
        dots_count = max(2, 70 - len(title))
        r_dots = p_t.add_run(" " + "." * dots_count + " ")
        r_dots.font.name = 'Times New Roman'
        r_dots.font.size = Pt(12)
        r_dots.font.color.rgb = RGBColor(120, 120, 120)
        
        r_page = p_t.add_run(page_num)
        r_page.font.name = 'Times New Roman'
        r_page.font.size = Pt(13)
        r_page.font.bold = is_bold

    doc.add_page_break()

    # ── 4. DANH MỤC TỪ VIẾT TẮT ─────────────────────────────────
    p_abbr_head = doc.add_paragraph()
    p_abbr_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abbr_head.paragraph_format.space_before = Pt(10)
    p_abbr_head.paragraph_format.space_after = Pt(12)
    p_abbr_head.paragraph_format.first_line_indent = Cm(0)
    r_abbr = p_abbr_head.add_run("DANH MỤC CÁC KÝ HIỆU VÀ CHỮ VIẾT TẮT")
    r_abbr.font.name = 'Times New Roman'
    r_abbr.font.size = Pt(15)
    r_abbr.font.bold = True

    abbr_headers = ["STT", "Ký hiệu / Viết tắt", "Tiếng Anh đầy đủ", "Giải thích ngữ nghĩa"]
    abbr_data = [
        ["1", "ACID", "Atomicity, Consistency, Isolation, Durability", "Các thuộc tính bảo đảm tính toàn vẹn của giao dịch cơ sở dữ liệu"],
        ["2", "API", "Application Programming Interface", "Giao diện lập trình ứng dụng"],
        ["3", "CLO", "Course Learning Outcomes", "Chuẩn đầu ra học phần / Chuẩn đánh giá đồ án"],
        ["4", "COD", "Cash On Delivery", "Phương thức thanh toán bằng tiền mặt khi nhận hàng"],
        ["5", "CRUD", "Create, Read, Update, Delete", "Bốn thao tác dữ liệu cơ bản: Tạo, Đọc, Sửa, Xóa"],
        ["6", "CSRF", "Cross-Site Request Forgery", "Tấn công giả mạo yêu cầu từ phía người dùng tin cậy"],
        ["7", "DOM", "Document Object Model", "Mô hình đối tượng tài liệu giao diện web"],
        ["8", "EJS", "Embedded JavaScript Templates", "Công cụ kết xuất mẫu HTML phía máy chủ cho Node.js"],
        ["9", "JSON", "JavaScript Object Notation", "Định dạng trao đổi dữ liệu dạng chuỗi đối tượng nhẹ"],
        ["10", "JWT", "JSON Web Token", "Chuẩn mã hóa thông tin an toàn giữa các bên dưới dạng JSON"],
        ["11", "MVC", "Model - View - Controller", "Mô hình kiến trúc phân tách Dữ liệu - Giao diện - Điều khiển"],
        ["12", "NoSQL", "Not Only SQL", "Hệ cơ sở dữ liệu phi quan hệ, lưu trữ linh hoạt dạng Document"],
        ["13", "ODM", "Object Data Modeling", "Mô hình hóa dữ liệu đối tượng cho cơ sở dữ liệu NoSQL"],
        ["14", "RBAC", "Role-Based Access Control", "Mô hình kiểm soát truy cập dựa trên vai trò người dùng"],
        ["15", "REST", "Representational State Transfer", "Kiểu kiến trúc thiết kế dịch vụ web chuẩn không trạng thái"],
        ["16", "SKU", "Stock Keeping Unit", "Mã đơn vị lưu kho duy nhất định danh từng phiên bản sản phẩm"],
        ["17", "SPA", "Single Page Application", "Ứng dụng web trang đơn, tải động không reload trang"],
        ["18", "SSR", "Server-Side Rendering", "Cơ chế tạo mã HTML hoàn chỉnh tại máy chủ trước khi gửi về client"],
        ["19", "UI/UX", "User Interface / User Experience", "Giao diện người dùng và Trải nghiệm người dùng"],
        ["20", "Vite", "Vite (tiếng Pháp: Nhanh)", "Công cụ đóng gói và phát triển frontend thế hệ mới tối ưu tốc độ"]
    ]
    add_table(doc, abbr_headers, abbr_data, [1.5, 3.5, 5.0, 6.0])

    # ── 5. DANH MỤC BẢNG BIỂU & HÌNH VẼ ────────────────────────
    p_lst_head = doc.add_paragraph()
    p_lst_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lst_head.paragraph_format.space_before = Pt(12)
    p_lst_head.paragraph_format.space_after = Pt(12)
    p_lst_head.paragraph_format.first_line_indent = Cm(0)
    r_lst = p_lst_head.add_run("DANH MỤC BẢNG BIỂU VÀ HÌNH VẼ")
    r_lst.font.name = 'Times New Roman'
    r_lst.font.size = Pt(15)
    r_lst.font.bold = True

    add_body_p(doc, "Bảng danh mục các bảng biểu số liệu trong báo cáo:", bold_prefix="A. Danh mục Bảng biểu: ", indent=False)
    table_items = [
        ("Bảng 1.1", "Phân tích các yêu cầu dữ liệu vào và dữ liệu ra của các phân hệ"),
        ("Bảng 1.2", "Tổng hợp các công nghệ và thư viện phần mềm sử dụng trong dự án"),
        ("Bảng 2.1", "Bảng phân tích yêu cầu chức năng hệ thống (Functional Requirements)"),
        ("Bảng 2.2", "Bảng phân tích yêu cầu phi chức năng hệ thống (Non-functional Requirements)"),
        ("Bảng 2.3", "Đặc tả chi tiết Use Case Đặt hàng và xử lý thanh toán (UC-04)"),
        ("Bảng 2.4", "Đặc tả chi tiết Schema thực thể Product trong MongoDB"),
        ("Bảng 2.5", "Đặc tả chi tiết Schema thực thể Order trong MongoDB"),
        ("Bảng 2.6", "Đặc tả chi tiết Schema thực thể User trong MongoDB"),
        ("Bảng 2.7", "Danh mục tổng hợp các RESTful API chính của hệ thống Rolex Boutique"),
        ("Bảng 3.1", "Cấu hình môi trường phần cứng và phần mềm triển khai thực nghiệm"),
        ("Bảng 3.2", "Bộ dữ liệu danh mục mẫu 31 mẫu đồng hồ tiêu biểu được khởi tạo"),
        ("Bảng 3.3", "Bảng ma trận các Test Case kiểm thử chức năng cốt lõi của hệ thống"),
        ("Bảng 3.4", "Thống kê các lỗi kỹ thuật phát sinh và giải pháp khắc phục thực tế"),
        ("Bảng 3.5", "So sánh định lượng hiệu quả trước và sau khi áp dụng 2 cải tiến kỹ thuật"),
        ("Bảng 3.6", "Bảng đối chiếu mức độ hoàn thành các mục tiêu đồ án ban đầu"),
        ("Bảng CLO", "Bảng ma trận đối chiếu chuẩn đầu ra học phần (Course Learning Outcomes)")
    ]
    for b_code, b_name in table_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.first_line_indent = Cm(0)
        r_c = p_item.add_run(f"{b_code}: ")
        r_c.font.bold = True
        p_item.add_run(b_name)

    add_body_p(doc, "Danh mục các hình vẽ, sơ đồ kiến trúc và ảnh chụp minh chứng trong báo cáo:", bold_prefix="B. Danh mục Hình vẽ và Sơ đồ: ", indent=False)
    fig_items = [
        ("Hình 2.1", "Sơ đồ Use Case tổng quát hệ thống website kinh doanh Rolex Boutique"),
        ("Hình 2.2", "Sơ đồ hoạt động (Activity Diagram) luồng Đặt hàng và xử lý tồn kho"),
        ("Hình 2.3", "Sơ đồ tuần tự (Sequence Diagram) luồng xác thực phiên và phân quyền quản trị"),
        ("Hình 2.4", "Sơ đồ lớp phân tích (Class Diagram) các thực thể nghiệp vụ cốt lõi"),
        ("Hình 2.5", "Sơ đồ mô hình dữ liệu quan hệ Document trong MongoDB"),
        ("Hình 2.6", "Bản thiết kế giao diện người dùng trên Figma (Wireframe/Mockup)"),
        ("Hình 3.1", "Giao diện trang chủ và sản phẩm vận hành thực tế trên môi trường máy chủ"),
        ("Hình 3.2", "Giao diện Bảng điều khiển quản trị hiện đại React SPA với biểu đồ Recharts"),
        ("Hình 3.3", "Cơ chế hoàn tác bù trừ (Compensating Rollback) khi đặt hàng nhiều sản phẩm")
    ]
    for f_code, f_name in fig_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.first_line_indent = Cm(0)
        r_c = p_item.add_run(f"{f_code}: ")
        r_c.font.bold = True
        p_item.add_run(f_name)

    # ── 6. MA TRẬN ĐỐI CHIẾU CHUẨN ĐẦU RA (CLO) ─────────────────
    p_clo_head = doc.add_paragraph()
    p_clo_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_clo_head.paragraph_format.space_before = Pt(14)
    p_clo_head.paragraph_format.space_after = Pt(8)
    p_clo_head.paragraph_format.first_line_indent = Cm(0)
    r_clo = p_clo_head.add_run("BẢNG ĐỐI CHIẾU CHUẨN ĐẦU RA (CLO MATRIX)\nKhoa Công nghệ số & Truyền thông – Trường Đại học Phương Đông")
    r_clo.font.name = 'Times New Roman'
    r_clo.font.size = Pt(14)
    r_clo.font.bold = True

    clo_headers = ["Mã CLO", "Tiêu chí đánh giá", "Tỷ lệ", "Vị trí trong báo cáo", "Minh chứng cụ thể trong đồ án"]
    clo_data = [
        ["CLO1", "Thiết kế & Báo cáo kỹ thuật", "30%", "Mở đầu; Chương 1; Chương 2; toàn bộ hình thức", "Bố cục học thuật, xác định rõ bài toán/phạm vi, trích dẫn chuẩn; Use case, Sequence diagram, ERD 5 collection, kiến trúc 3 tầng lai, 25+ RESTful API."],
        ["CLO2", "Sản phẩm phần mềm thực thi", "30%", "Mục 3.1, Mục 3.2; Demo thực tế sản phẩm", "Ứng dụng chạy hoàn chỉnh không lỗi trên Node.js/MongoDB; giỏ hàng, đặt hàng, quản trị kho, lịch sử đơn; 31 mẫu đồng hồ seed data; 25 test case có kết quả thực tế."],
        ["CLO3", "Giải pháp cải tiến kỹ thuật", "20%", "Mục 3.3; Bằng chứng so sánh trước/sau", "Cải tiến 1: Khấu trừ tồn kho nguyên tử kết hợp Hoàn tác bù trừ (Compensating Rollback) trên MongoDB Standalone. Cải tiến 2: Nâng cấp Admin từ EJS sang React 18/Vite SPA + Zustand + Recharts."],
        ["CLO4", "Tự học & Định hướng phát triển", "20%", "Mục 3.4; Phụ lục & Kế hoạch cá nhân", "Tự học chuyên sâu React 18, TypeScript, Tailwind, Zustand, Mongoose Aggregation Pipeline; lập kế hoạch phát triển Microservices, Docker và tích hợp cổng thanh toán Sandbox."]
    ]
    add_table(doc, clo_headers, clo_data, [1.8, 3.2, 1.2, 4.0, 5.8])

    doc.add_page_break()
